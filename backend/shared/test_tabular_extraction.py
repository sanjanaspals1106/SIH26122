"""
Unit tests for shared/tabular_extraction.py -- CSV/XLSX reading, activity-id
column detection, and structured row-to-claim construction. Includes both
hand-built fixtures and the real progress-report-csv / discipline-report-
xlsx sample_data fixtures.
"""
import io
from pathlib import Path

import openpyxl
import pytest

from backend.shared.tabular_extraction import (
    build_claim_from_row,
    detect_progress_columns,
    normalize_pct,
    read_tabular_file,
    to_float,
)

_SAMPLE_ROOT = Path(__file__).resolve().parents[2] / "sample_data" / "input"
_REAL_CSV = _SAMPLE_ROOT / "progress-report-csv" / "daily_progress_2026-08-14.csv"
_REAL_XLSX = _SAMPLE_ROOT / "discipline-report-xlsx" / "discipline_progress_2026-08-14.xlsx"


def test_detect_progress_columns_prefers_activity_id_over_task_id():
    """This project's real progress-report fixtures use "Task ID"/"L6 Task
    ID" for a per-report SUB-task suffix (e.g. "CIV-PS3-TR-0180-01"), not
    the schedule's real activity_id -- must never be picked over a genuine
    Activity ID/L5 Activity ID column when both are present."""
    cols = detect_progress_columns(["Activity ID", "Task ID", "Discipline", "Today Actual"])
    assert cols is not None
    assert cols.activity_col == "Activity ID"


def test_detect_progress_columns_l5_variant():
    cols = detect_progress_columns(["L5 Activity ID", "L6 Task ID", "Work description"])
    assert cols is not None
    assert cols.activity_col == "L5 Activity ID"


def test_detect_progress_columns_none_without_activity_id():
    """A "Task ID"-only sheet (no genuine activity-id column) must NOT be
    treated as structured -- callers fall back to LLM extraction instead of
    silently attaching claims to the wrong id."""
    cols = detect_progress_columns(["Task ID", "Description", "Notes"])
    assert cols is None


def test_normalize_pct_converts_excel_fraction():
    """Excel stores a percentage-formatted cell as its raw 0-1 fraction
    (0.917 for a displayed 91.7%), not the 0-100 scale this project's CSV
    exports otherwise use -- must be normalized to 91.7, not passed through
    as literally "0.917%"."""
    assert normalize_pct(0.917) == pytest.approx(91.7)
    assert normalize_pct(1.0) == 100.0
    assert normalize_pct(100.0) == 100.0  # already on 0-100 scale, unchanged
    assert normalize_pct(None) is None
    assert normalize_pct("") is None


def test_to_float_handles_nan_and_blank():
    assert to_float("nan") is None
    assert to_float("") is None
    assert to_float(None) is None
    assert to_float("40") == 40.0


def test_read_tabular_file_csv_bom_and_semicolon():
    csv_text = "Activity ID;Today Actual;Progress Pct\nCIV-A;40;100.0\n"
    contents = b"\xef\xbb\xbf" + csv_text.encode("utf-8")
    sheets = read_tabular_file("x.csv", contents)
    assert list(sheets.keys()) == [""]
    df = sheets[""]
    assert list(df.columns) == ["Activity ID", "Today Actual", "Progress Pct"]
    assert df.iloc[0]["Activity ID"] == "CIV-A"


def test_read_tabular_file_csv_empty_raises():
    with pytest.raises(ValueError):
        read_tabular_file("x.csv", b"")


def test_read_tabular_file_corrupt_xlsx_raises():
    with pytest.raises(ValueError):
        read_tabular_file("x.xlsx", b"not a real xlsx file")


def test_read_tabular_file_xlsx_reads_every_sheet():
    """The actual root cause of the original Excel intake bug:
    pandas.read_excel defaults to sheet_name=0 (first sheet only). A
    3-sheet workbook must yield all 3 sheets, not 1."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, activity in [("Civil", "CIV-A"), ("Piping", "PIP-A"), ("Electrical", "ELE-A")]:
        ws = wb.create_sheet(name)
        ws.append(["Activity ID", "Today Actual", "Progress Pct"])
        ws.append([activity, 10, 50.0])
    buf = io.BytesIO()
    wb.save(buf)
    sheets = read_tabular_file("x.xlsx", buf.getvalue())
    assert set(sheets.keys()) == {"Civil", "Piping", "Electrical"}


def test_read_tabular_file_xlsx_title_row_before_header():
    """Some real exports put a title/banner row before the real header
    (see discipline-report-xlsx fixtures) -- the real header must still be
    found and used, not the title row misread as headers."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Civil Daily Discipline Report"])
    ws.append(["Some subtitle line"])
    ws.append([])
    ws.append(["Activity ID", "Today Actual", "Progress Pct"])
    ws.append(["CIV-A", 40, 100.0])
    buf = io.BytesIO()
    wb.save(buf)
    sheets = read_tabular_file("x.xlsx", buf.getvalue())
    df = sheets["Sheet"]
    cols = detect_progress_columns(list(df.columns))
    assert cols is not None
    assert cols.activity_col == "Activity ID"
    assert df.iloc[0]["Activity ID"] == "CIV-A"


def test_build_claim_from_row_skips_no_progress():
    cols = detect_progress_columns(["Activity ID", "Today Actual", "Progress Pct"])
    row = {"Activity ID": "CIV-A", "Today Actual": 0, "Progress Pct": 0.0}
    assert build_claim_from_row(row, cols) is None


def test_build_claim_from_row_skips_blank_activity_id():
    cols = detect_progress_columns(["Activity ID", "Today Actual", "Progress Pct"])
    row = {"Activity ID": "", "Today Actual": 40, "Progress Pct": 100.0}
    assert build_claim_from_row(row, cols) is None


def test_build_claim_from_row_pct_present_uses_cumulative_mode():
    cols = detect_progress_columns(["Activity ID", "Today Actual", "Progress Pct"])
    row = {"Activity ID": "CIV-A", "Today Actual": 40, "Progress Pct": 91.7}
    raw_text, extracted = build_claim_from_row(row, cols)
    assert extracted.claim_mode.value == "CUMULATIVE_PCT"
    assert extracted.claimed_pct == 91.7
    assert extracted.claimed_quantity is None


def test_build_claim_from_row_qty_only_uses_incremental_mode():
    """A row with a quantity column but no percentage column must not
    silently drop the quantity -- it becomes an INCREMENTAL_QUANTITY claim
    instead of a claim with every numeric field null."""
    cols = detect_progress_columns(["Activity ID", "Today Actual", "Unit"])
    row = {"Activity ID": "CIV-A", "Today Actual": 40, "Unit": "m"}
    raw_text, extracted = build_claim_from_row(row, cols)
    assert extracted.claim_mode.value == "INCREMENTAL_QUANTITY"
    assert extracted.claimed_quantity == 40.0
    assert extracted.claimed_uom == "m"
    assert extracted.claimed_pct is None


def test_build_claim_from_row_discipline_hint_fallback():
    """A sheet with no Discipline column (discipline implied by the sheet
    name itself, as in the real discipline-report-xlsx fixtures) must still
    resolve a real discipline via the caller-supplied hint."""
    cols = detect_progress_columns(["Activity ID", "Today Actual", "Progress Pct"])
    row = {"Activity ID": "CIV-A", "Today Actual": 40, "Progress Pct": 100.0}
    _, extracted = build_claim_from_row(row, cols, discipline_hint="Civil")
    assert extracted.discipline.value == "CIVIL"


def test_build_claim_from_row_unrecognized_discipline_stays_null():
    """An unrecognized discipline label must not be force-fit into an
    invalid enum value -- ExtractedClaimFields would reject it outright."""
    cols = detect_progress_columns(["Activity ID", "Discipline", "Today Actual", "Progress Pct"])
    row = {"Activity ID": "CIV-A", "Discipline": "Landscaping", "Today Actual": 40, "Progress Pct": 100.0}
    _, extracted = build_claim_from_row(row, cols)
    assert extracted.discipline is None


@pytest.mark.skipif(not _REAL_CSV.exists(), reason="real sample CSV not present")
def test_real_sample_csv_end_to_end():
    contents = _REAL_CSV.read_bytes()
    sheets = read_tabular_file(_REAL_CSV.name, contents)
    df = sheets[""]
    cols = detect_progress_columns(list(df.columns))
    assert cols is not None
    assert cols.activity_col == "Activity ID"

    claims = [build_claim_from_row(row.to_dict(), cols) for _, row in df.iterrows()]
    claims = [c for c in claims if c is not None]
    assert len(claims) == 4
    ids = {c[1].reported_activity_id for c in claims}
    assert ids == {"CIV-PS3-TR-0180", "PIP-PS3-WLD-024", "ELE-PS3-CT-011", "MECH-PS3-DWP-003"}


@pytest.mark.skipif(not _REAL_XLSX.exists(), reason="real sample XLSX not present")
def test_real_sample_xlsx_end_to_end():
    contents = _REAL_XLSX.read_bytes()
    sheets = read_tabular_file(_REAL_XLSX.name, contents)
    assert set(sheets.keys()) == {"Civil", "Piping", "Electrical"}

    all_claims = []
    for sheet_name, df in sheets.items():
        cols = detect_progress_columns(list(df.columns))
        assert cols is not None, f"sheet {sheet_name} should have a detectable activity-id column"
        assert cols.activity_col == "L5 Activity ID", "must prefer L5 Activity ID over the L6 Task ID sub-task column"
        for _, row in df.iterrows():
            result = build_claim_from_row(row.to_dict(), cols, discipline_hint=sheet_name)
            if result is not None:
                all_claims.append(result)

    assert len(all_claims) == 3
    by_activity = {extracted.reported_activity_id: extracted for _, extracted in all_claims}
    assert by_activity["CIV-PS3-TR-0180"].discipline.value == "CIVIL"
    assert by_activity["PIP-PS3-WLD-024"].discipline.value == "PIPING"
    assert by_activity["ELE-PS3-CT-011"].discipline.value == "ELECTRICAL"
    # The XLSX stores the exact fraction 22/24 = 0.91666..., while the CSV
    # sample rounds the same underlying value to "91.7" for display -- both
    # are correct for their own source, so compare loosely.
    assert by_activity["PIP-PS3-WLD-024"].claimed_pct == pytest.approx(91.7, abs=0.05)
