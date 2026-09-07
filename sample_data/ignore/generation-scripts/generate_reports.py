import os
import csv
import json
from PIL import Image, ImageDraw, ImageFont
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

BASE_DIR = r"D:\SIH26122\sample_data"
INPUT_DIR = os.path.join(BASE_DIR, "input")

TXT_DIR = os.path.join(INPUT_DIR, "daily-report-txt")
PDF_DIR = os.path.join(INPUT_DIR, "daily-report-pdf")
XLSX_DIR = os.path.join(INPUT_DIR, "discipline-report-xlsx")
CSV_DIR = os.path.join(INPUT_DIR, "progress-report-csv")
DIARY_DIR = os.path.join(INPUT_DIR, "scanned-diary")
JSON_DIR = os.path.join(INPUT_DIR, "field-reports-json")

for d in [TXT_DIR, PDF_DIR, XLSX_DIR, CSV_DIR, DIARY_DIR, JSON_DIR]:
    os.makedirs(d, exist_ok=True)

# -------------------------------------------------------------
# 1. DAILY PROGRESS REPORTS (TXT)
# -------------------------------------------------------------
TXT_REPORTS = {
    "2026-08-11": {
        "date_str": "11 Aug 2026",
        "weather": "Clear and sunny. 38 C. High humidity.",
        "manpower": 22,
        "equipment": "Excavator 1; Mobile Crane 1; Welding Generator 2",
        "supervisor": "A. Das, Site Supervisor",
        "items": [
            ("CIV-PS3-FND-001", "Civil", "Excavation and blinding concrete for pump base commenced.", "20 m3", "20/65 m3", "30.8%"),
            ("PIP-PS3-WLD-024", "Piping", "Utility header fit-up and first pass root welding started.", "6 joints", "6/24 joints", "25.0%"),
            ("PIP-PS3-SPO-015", "Piping", "Firewater underground spool placement in trench.", "20 m", "20/80 m", "25.0%"),
            ("ELE-PS3-CT-011", "Electrical", "Sieved sand bedding hauling and placement at MCC-02 trench.", "40 m", "60/160 m", "37.5%"),
            ("HSE-PS3-IND-001", "HSE", "Site induction and heat stress awareness toolbox talk conducted.", "1 day", "1/10 days", "10.0%")
        ],
        "quality": "Welding fit-up inspected and accepted for 6 joints. Sand gradation certificate checked.",
        "plan_next": "Continue pump base blinding and utility header hot pass welding."
    },
    "2026-08-13": {
        "date_str": "13 Aug 2026",
        "weather": "Partly cloudy. 35 C. Light breeze.",
        "manpower": 24,
        "equipment": "Excavator 1; Crane 1; Welding Genset 2; Compactor 1",
        "supervisor": "A. Das, Site Supervisor",
        "items": [
            ("CIV-PS3-FND-001", "Civil", "Blinding concrete pour completed for pump foundation. Final level checked.", "20 m3", "65/65 m3", "100.0%"),
            ("PIP-PS3-WLD-024", "Piping", "Six field weld joints completed on utility header. Visual testing clear.", "6 joints", "20/24 joints", "83.3%"),
            ("PIP-PS3-SPO-015", "Piping", "Firewater spool alignment and tacking completed.", "20 m", "65/80 m", "81.2%"),
            ("ELE-PS3-CT-011", "Electrical", "Cable trench bedding placement continued at MCC-02.", "30 m", "120/160 m", "75.0%"),
            ("HSE-PS3-IND-001", "HSE", "Daily safety briefing and trench collapse hazard review.", "1 day", "3/10 days", "30.0%")
        ],
        "quality": "Concrete blinding test cubes taken. Pipe weld visual inspection 100% accepted.",
        "plan_next": "Begin rebar fixing for pump foundation; prepare for utility trench excavation at CH 0+180."
    },
    "2026-08-15": {
        "date_str": "15 Aug 2026",
        "weather": "Dry and clear. 36 C. Light wind.",
        "manpower": 26,
        "equipment": "Excavator 1; Mobile Crane 1; Air Compressor 1; Welding Rig 1",
        "supervisor": "A. Das, Site Supervisor",
        "items": [
            ("CIV-PS3-TR-0220", "Civil", "Trench excavation CH 0+220 to 0+260 commenced.", "20 m", "20/40 m", "50.0%"),
            ("PIP-PS3-WLD-024", "Piping", "Final two weld joints completed. Header welding completed 100%.", "2 joints", "24/24 joints", "100.0%"),
            ("ELE-PS3-CT-011", "Electrical", "Final 20m cable trench bedding placed at MCC-02. Ready for cable pull.", "20 m", "160/160 m", "100.0%"),
            ("EQP-PS3-TK-001", "Static/Rotating Equipment", "Sump tank TK-01 internal inspection and nozzle orientation check accepted.", "1 ea", "1/1 ea", "100.0%"),
            ("HSE-PS3-AUD-001", "HSE", "Weekly environmental compliance audit conducted with client representative.", "1 audit", "1/2 audits", "50.0%")
        ],
        "quality": "Utility header radiographic testing films reviewed; zero defects reported.",
        "plan_next": "Continue CH 0+220 trench excavation and prepare pump base rebar inspection."
    },
    "2026-08-19": {
        "date_str": "19 Aug 2026",
        "weather": "Hot and sunny. 39 C. Moderate dust.",
        "manpower": 28,
        "equipment": "Excavator 1; Heavy Crane 50T 1; Concrete Transit Mixer 2; Boom Pump 1",
        "supervisor": "K. Sharma, Construction Supervisor",
        "items": [
            ("CIV-PS3-FND-003", "Civil", "Structural concrete pour for Pump P-101/P-102 foundation started.", "30 m3", "30/60 m3", "50.0%"),
            ("PIP-PS3-HDR-100-B", "Piping", "Utility header Section B spool placement and bolt-up in progress.", "8 m", "18/30 m", "60.0%"),
            ("PIP-PS3-HYD-001", "Piping", "Hydrostatic pressure testing of Firewater Loop Sector 1 completed at 24 bar.", "1 test", "1/1 test", "100.0%"),
            ("EQP-PS3-PMP-101", "Static/Rotating Equipment", "Epoxy grouting of baseplate for Crude Transfer Pump P-101 completed.", "1 ea", "1/1 ea", "100.0%"),
            ("INS-PS3-PT-021", "Instrumentation", "Pressure transmitter PT-1021 stanchion mounting and impulse tubing.", "2 ea", "2/6 ea", "33.3%")
        ],
        "quality": "Hydrotest pressure held for 4 hours with no pressure drop. Concrete slump measured at 130mm.",
        "plan_next": "Complete pump foundation pour and proceed with Crude Transfer Pump P-102 rigging."
    },
    "2026-08-21": {
        "date_str": "21 Aug 2026",
        "weather": "Humid and overcast. 34 C. Wind 15 km/h.",
        "manpower": 25,
        "equipment": "Crane 50T 1; Rigging truck 1; Pipe support welding rig 1",
        "supervisor": "K. Sharma, Construction Supervisor",
        "items": [
            ("CIV-PS3-FND-003", "Civil", "Final concrete pour completed for pump base. Curing started.", "30 m3", "60/60 m3", "100.0%"),
            ("EQP-PS3-PMP-102", "Static/Rotating Equipment", "Crude Transfer Pump P-102 positioning on plinth and laser alignment accepted.", "1 ea", "1/1 ea", "100.0%"),
            ("PIP-PS3-TIE-002", "Piping", "Tie-in spool fit-up at existing manifold M-01 commenced.", "2 joints", "2/4 joints", "50.0%"),
            ("INS-PS3-CAL-003", "Instrumentation", "Bench calibration of control valves FCV-101 to 104 completed.", "4 ea", "4/10 ea", "40.0%"),
            ("HSE-PS3-GAS-001", "HSE", "Confined space atmospheric testing for tie-in manifold pit.", "6 checks", "32/40 checks", "80.0%")
        ],
        "quality": "Laser alignment radial/axial runout verified < 0.04mm. Gas check O2 20.9%, LEL 0%.",
        "plan_next": "Torque manifold tie-in flanges and continue valve bench calibration."
    }
}

def generate_txt_reports():
    for dt, data in TXT_REPORTS.items():
        fname = os.path.join(TXT_DIR, f"daily_progress_report_{dt}.txt")
        if os.path.exists(fname) and dt == "2026-08-14":
            continue
        lines = [
            "DAILY PROGRESS REPORT",
            "Project: North Field Utility Corridor",
            f"Report date: {data['date_str']}",
            "Location: Pump Station 3",
            "Reporting period: 0800 h to 1800 h",
            "",
            "WEATHER",
            data["weather"],
            "",
            "WORK EXECUTED"
        ]
        for idx, item in enumerate(data["items"], 1):
            lines.append(f"{idx}. {item[0]} | {item[1]} | {item[2]}")
            lines.append(f"   Today: {item[3]} | Cumulative: {item[4]} | Physical progress: {item[5]}")
        lines.extend([
            "",
            "QUALITY AND HOLD POINTS",
            f"- {data['quality']}",
            "",
            "RESOURCES",
            f"Manpower: {data['manpower']}",
            f"Plant and equipment: {data['equipment']}",
            "",
            "PLAN FOR NEXT SHIFT",
            f"- {data['plan_next']}",
            "",
            f"Prepared by: {data['supervisor']}",
            "Fictional demonstration data for SIH26122 format testing.",
            ""
        ])
        with open(fname, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
        print(f"Generated TXT DPR: {fname}")

# -------------------------------------------------------------
# 2. DAILY PROGRESS REPORTS (PDF)
# -------------------------------------------------------------
PDF_REPORTS = {
    "2026-08-12": {
        "date_str": "12 Aug 2026",
        "weather": "Sunny. 37 C.",
        "manpower": 23,
        "items": [
            ("CIV-PS3-FND-001", "Civil", "Blinding concrete pour continued for pump foundation.", "45/65 m3", "69.2%", "Ongoing"),
            ("PIP-PS3-WLD-024", "Piping", "Field weld joints root and hot passes accepted.", "14/24 joints", "58.3%", "Ongoing"),
            ("PIP-PS3-SPO-015", "Piping", "Firewater spool trench placement.", "45/80 m", "56.2%", "Ongoing"),
            ("ELE-PS3-CT-011", "Electrical", "Cable trench bedding placement at MCC-02.", "90/160 m", "56.2%", "Ongoing"),
            ("INS-PS3-JB-001", "Instrumentation", "Junction box stanchions and JB-101 mounting.", "4/12 ea", "33.3%", "Ongoing")
        ],
        "quality": "Visual inspection accepted on welded joints. Trench bedding depth verified 100mm.",
        "next_shift": "Continue welding on utility header and complete foundation blinding."
    },
    "2026-08-18": {
        "date_str": "18 Aug 2026",
        "weather": "Overcast, slight haze. 36 C.",
        "manpower": 27,
        "items": [
            ("CIV-PS3-FND-002", "Civil", "Pump P-101/P-102 foundation rebar and shuttering completed.", "12/12 t", "100.0%", "Complete"),
            ("CIV-PS3-TR-0220", "Civil", "Utility trench excavation CH 0+220 to 0+260 completed.", "40/40 m", "100.0%", "Complete"),
            ("PIP-PS3-HDR-100-B", "Piping", "Utility header spool Section B fit-up in progress.", "10/30 m", "33.3%", "Ongoing"),
            ("EQP-PS3-PMP-101", "Static/Rotating Equipment", "Crude Transfer Pump P-101 foundation prep and shimming.", "0/1 ea", "0.0%", "Ongoing"),
            ("ELE-PS3-TR-005", "Electrical", "Cable tray installation in Main Piperack Tier 2.", "90/200 m", "45.0%", "Ongoing")
        ],
        "quality": "Rebar spacing and concrete cover blocks inspected and approved prior to pour clearance.",
        "next_shift": "Pre-pour inspection for pump foundation; continue Tier 2 cable tray installation."
    },
    "2026-08-21": {
        "date_str": "21 Aug 2026",
        "weather": "Humid and cloudy. 34 C.",
        "manpower": 25,
        "items": [
            ("CIV-PS3-FND-003", "Civil", "Pump foundation structural concrete pour finished.", "60/60 m3", "100.0%", "Complete"),
            ("EQP-PS3-PMP-102", "Static/Rotating Equipment", "Crude Transfer Pump P-102 laser shaft alignment completed.", "1/1 ea", "100.0%", "Complete"),
            ("PIP-PS3-TIE-002", "Piping", "Tie-in spool fit-up at existing manifold M-01.", "2/4 joints", "50.0%", "Ongoing"),
            ("INS-PS3-CAL-003", "Instrumentation", "Control valve bench calibration at workshop.", "4/10 ea", "40.0%", "Ongoing"),
            ("HSE-PS3-GAS-001", "HSE", "Confined space atmospheric gas testing.", "32/40 checks", "80.0%", "Ongoing")
        ],
        "quality": "Pump laser alignment within 0.03mm tolerance. Gas testing zero LEL confirmed.",
        "next_shift": "Complete manifold bolt torqueing and continue control valve stroke testing."
    }
}

def generate_pdf_reports():
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=15,
        leading=18,
        textColor=colors.HexColor("#1A365D")
    )
    sub_style = ParagraphStyle(
        "ReportSub",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#4A5568")
    )
    head_style = ParagraphStyle(
        "SectionHead",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#2B6CB0"),
        spaceBefore=6,
        spaceAfter=4
    )
    cell_style = ParagraphStyle(
        "CellText",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=10
    )
    cell_bold = ParagraphStyle(
        "CellBold",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10
    )

    for dt, data in PDF_REPORTS.items():
        fname = os.path.join(PDF_DIR, f"daily_progress_report_{dt}.pdf")
        if os.path.exists(fname) and dt == "2026-08-14":
            continue

        doc = SimpleDocTemplate(
            fname,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )
        elements = []

        elements.append(Paragraph("Daily Progress Report", title_style))
        elements.append(Paragraph(f"North Field Utility Corridor | Pump Station 3 Tie In | {data['date_str']}", sub_style))
        elements.append(Spacer(1, 8))

        # Metadata Table
        meta_data = [
            [Paragraph("<b>Report date:</b> " + data["date_str"], cell_style), Paragraph("<b>Reporting period:</b> 0800 h to 1800 h", cell_style)],
            [Paragraph("<b>Weather:</b> " + data["weather"], cell_style), Paragraph(f"<b>Manpower:</b> {data['manpower']} personnel on site", cell_style)]
        ]
        meta_table = Table(meta_data, colWidths=[270, 270])
        meta_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EDF2F7")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E0")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elements.append(meta_table)
        elements.append(Spacer(1, 10))

        # Progress Section
        elements.append(Paragraph("Progress against linked schedule activities", head_style))
        table_data = [[
            Paragraph("Activity ID", cell_bold),
            Paragraph("Discipline", cell_bold),
            Paragraph("Work completed today", cell_bold),
            Paragraph("Actual / Plan", cell_bold),
            Paragraph("Progress", cell_bold),
            Paragraph("Status", cell_bold)
        ]]

        for item in data["items"]:
            table_data.append([
                Paragraph(item[0], cell_style),
                Paragraph(item[1], cell_style),
                Paragraph(item[2], cell_style),
                Paragraph(item[3], cell_style),
                Paragraph(item[4], cell_style),
                Paragraph(item[5], cell_style)
            ])

        prog_table = Table(table_data, colWidths=[90, 70, 190, 75, 55, 60])
        prog_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2B6CB0")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E0")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7FAFC")]),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elements.append(prog_table)
        elements.append(Spacer(1, 10))

        # Quality and look-ahead
        elements.append(Paragraph("Quality, constraints and next shift", head_style))
        q_data = [
            [Paragraph("<b>Quality evidence:</b>", cell_style), Paragraph(data["quality"], cell_style)],
            [Paragraph("<b>Next shift:</b>", cell_style), Paragraph(data["next_shift"], cell_style)]
        ]
        q_table = Table(q_data, colWidths=[100, 440])
        q_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        elements.append(q_table)
        elements.append(Spacer(1, 12))

        elements.append(Paragraph("Prepared by: Site Construction Supervisor | Fictional demonstration data for SIH26122 format testing.", sub_style))

        doc.build(elements)
        print(f"Generated PDF DPR: {fname}")

# -------------------------------------------------------------
# 3. DISCIPLINE REPORTS (XLSX)
# -------------------------------------------------------------
XLSX_REPORTS = {
    "2026-08-18": {
        "date_str": "18 Aug 2026",
        "sheets": {
            "Civil": [
                ("CIV-PS3-FND-002", "CIV-PS3-FND-002-01", "Pump P-101/P-102 foundation rebar and shuttering", "t", 12.0, 4.0, 8.0, 12.0, "Rebar fixing inspected and approved"),
                ("CIV-PS3-TR-0220", "CIV-PS3-TR-0220-01", "Utility trench excavation CH 0+220 to 0+260", "m", 40.0, 20.0, 20.0, 40.0, "Trench depth verified along chainage")
            ],
            "Piping": [
                ("PIP-PS3-HDR-100-B", "PIP-PS3-HDR-100-B-01", "Utility header spool Section B fit-up and erection", "m", 30.0, 0.0, 10.0, 10.0, "Fit-up inspection passed")
            ],
            "Electrical": [
                ("ELE-PS3-TR-005", "ELE-PS3-TR-005-01", "Cable tray installation Main Piperack Tier 2", "m", 200.0, 40.0, 50.0, 90.0, "Tray alignment and grounding jumper verified")
            ]
        }
    },
    "2026-08-20": {
        "date_str": "20 Aug 2026",
        "sheets": {
            "Civil": [
                ("CIV-PS3-DRN-001", "CIV-PS3-DRN-001-01", "Stormwater drainage channel installation Section A", "m", 120.0, 90.0, 20.0, 110.0, "Invert elevation checked against benchmark")
            ],
            "Piping": [
                ("PIP-PS3-HDR-100-B", "PIP-PS3-HDR-100-B-01", "Utility header spool Section B fit-up and erection", "m", 30.0, 18.0, 7.0, 25.0, "Flange alignment accepted")
            ],
            "Electrical": [
                # EDGE-004: In XLSX this is reported as 140 m / 200 m = 70.0% (whereas 20-Aug CSV reports 120 m / 200 m = 60.0%)
                ("ELE-PS3-TR-005", "ELE-PS3-TR-005-01", "Cable tray installation Main Piperack Tier 2", "m", 200.0, 90.0, 50.0, 140.0, "Tier 2 tray run complete to grid line 4")
            ]
        }
    }
}

def generate_xlsx_reports():
    header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1A365D")
    sub_font = Font(name="Calibri", size=10, italic=True, color="4A5568")
    bold_font = Font(name="Calibri", size=10, bold=True)
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E0'),
        right=Side(style='thin', color='CBD5E0'),
        top=Side(style='thin', color='CBD5E0'),
        bottom=Side(style='thin', color='CBD5E0')
    )

    for dt, data in XLSX_REPORTS.items():
        fname = os.path.join(XLSX_DIR, f"discipline_progress_{dt}.xlsx")
        if os.path.exists(fname) and dt == "2026-08-14":
            continue

        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        for sheet_name, rows in data["sheets"].items():
            ws = wb.create_sheet(title=sheet_name)
            ws.views.sheetView[0].showGridLines = True

            ws.cell(row=2, column=1, value=f"{sheet_name} Daily Discipline Report").font = title_font
            ws.cell(row=3, column=1, value=f"North Field Utility Corridor | Pump Station 3 Tie In | Report date: {data['date_str']}").font = sub_font

            headers = [
                "L5 Activity ID", "L6 Task ID", "Work description", "Unit",
                "Planned qty", "Prior actual", "Today actual", "Cumulative actual",
                "Progress", "Status and evidence"
            ]

            for col_num, h_text in enumerate(headers, 1):
                c = ws.cell(row=5, column=col_num, value=h_text)
                c.fill = header_fill
                c.font = header_font
                c.alignment = Alignment(horizontal="center", vertical="center")

            current_row = 6
            for r in rows:
                ws.cell(row=current_row, column=1, value=r[0]).font = bold_font
                ws.cell(row=current_row, column=2, value=r[1]).font = regular_font
                ws.cell(row=current_row, column=3, value=r[2]).font = regular_font
                ws.cell(row=current_row, column=4, value=r[3]).font = regular_font
                ws.cell(row=current_row, column=5, value=r[4]).font = regular_font
                ws.cell(row=current_row, column=6, value=r[5]).font = regular_font
                ws.cell(row=current_row, column=7, value=r[6]).font = regular_font
                ws.cell(row=current_row, column=8, value=r[7]).font = bold_font
                # Formula for progress: Cumulative / Planned
                prog_cell = ws.cell(row=current_row, column=9, value=f"=H{current_row}/E{current_row}")
                prog_cell.number_format = '0.0%'
                prog_cell.font = bold_font
                ws.cell(row=current_row, column=10, value=r[8]).font = regular_font

                for c_idx in range(1, 11):
                    ws.cell(row=current_row, column=c_idx).border = thin_border
                current_row += 1

            # Summary section
            summary_row = current_row + 2
            ws.cell(row=summary_row, column=1, value="Report summary").font = bold_font
            ws.cell(row=summary_row, column=2, value="Value").font = bold_font
            ws.cell(row=summary_row, column=4, value="Site conditions").font = bold_font
            ws.cell(row=summary_row, column=5, value="Value").font = bold_font

            ws.cell(row=summary_row + 1, column=1, value="Total activities").font = regular_font
            ws.cell(row=summary_row + 1, column=2, value=len(rows)).font = bold_font
            ws.cell(row=summary_row + 1, column=4, value="Weather").font = regular_font
            ws.cell(row=summary_row + 1, column=5, value="Good working conditions").font = regular_font

            ws.column_dimensions['A'].width = 18
            ws.column_dimensions['B'].width = 20
            ws.column_dimensions['C'].width = 38
            ws.column_dimensions['D'].width = 8
            ws.column_dimensions['E'].width = 12
            ws.column_dimensions['F'].width = 12
            ws.column_dimensions['G'].width = 12
            ws.column_dimensions['H'].width = 16
            ws.column_dimensions['I'].width = 12
            ws.column_dimensions['J'].width = 35

        wb.save(fname)
        print(f"Generated XLSX Discipline Report: {fname}")

# -------------------------------------------------------------
# 4. PROGRESS REPORTS (CSV)
# -------------------------------------------------------------
CSV_PROGRESS_DATA = {
    "2026-08-11": [
        ("CIV-PS3-FND-001", "CIV-PS3-FND-001-01", "Civil", "Pump foundation excavation and blinding", "m3", 65.0, 0.0, 20.0, 20.0, 30.8, "Ongoing", "daily_progress_report_2026-08-11.txt", "Blinding pour in progress"),
        ("PIP-PS3-WLD-024", "PIP-PS3-WLD-024-01", "Piping", "Field weld joints utility header", "joints", 24.0, 0.0, 6.0, 6.0, 25.0, "Ongoing", "daily_progress_report_2026-08-11.txt", "Root pass welding started"),
        ("PIP-PS3-SPO-015", "PIP-PS3-SPO-015-01", "Piping", "Firewater underground spool installation", "m", 80.0, 0.0, 20.0, 20.0, 25.0, "Ongoing", "daily_progress_report_2026-08-11.txt", "Spools strung in trench"),
        ("ELE-PS3-CT-011", "ELE-PS3-CT-011-02", "Electrical", "Place cable trench bedding at MCC-02", "m", 160.0, 20.0, 40.0, 60.0, 37.5, "Ongoing", "daily_progress_report_2026-08-11.txt", "Sand bedding placement"),
        ("HSE-PS3-IND-001", "HSE-PS3-IND-001-01", "HSE", "Site safety induction and toolbox talk", "days", 10.0, 0.0, 1.0, 1.0, 10.0, "Ongoing", "daily_progress_report_2026-08-11.txt", "Toolbox completed")
    ],
    "2026-08-12": [
        ("CIV-PS3-FND-001", "CIV-PS3-FND-001-01", "Civil", "Pump foundation excavation and blinding", "m3", 65.0, 20.0, 25.0, 45.0, 69.2, "Ongoing", "daily_progress_report_2026-08-12.pdf", "Blinding concrete continuous pour"),
        ("PIP-PS3-WLD-024", "PIP-PS3-WLD-024-01", "Piping", "Field weld joints utility header", "joints", 24.0, 6.0, 8.0, 14.0, 58.3, "Ongoing", "daily_progress_report_2026-08-12.pdf", "8 joints hot pass accepted"),
        ("PIP-PS3-HDR-100-A", "PIP-PS3-HDR-100-A-01", "Piping", "Utility header spool Section A", "m", 40.0, 0.0, 15.0, 15.0, 37.5, "Ongoing", "daily_progress_report_2026-08-12.pdf", "Section A placement"),
        ("PIP-PS3-SPO-015", "PIP-PS3-SPO-015-01", "Piping", "Firewater underground spool installation", "m", 80.0, 20.0, 25.0, 45.0, 56.2, "Ongoing", "daily_progress_report_2026-08-12.pdf", "Trench pipe alignment"),
        ("ELE-PS3-CT-011", "ELE-PS3-CT-011-02", "Electrical", "Place cable trench bedding at MCC-02", "m", 160.0, 60.0, 30.0, 90.0, 56.2, "Ongoing", "daily_progress_report_2026-08-12.pdf", "Trench bedding progressing"),
        ("INS-PS3-JB-001", "INS-PS3-JB-001-01", "Instrumentation", "Field junction box installation", "ea", 12.0, 0.0, 4.0, 4.0, 33.3, "Ongoing", "daily_progress_report_2026-08-12.pdf", "JB-101 to 104 mounted"),
        ("HSE-PS3-IND-001", "HSE-PS3-IND-001-01", "HSE", "Site safety induction and toolbox talk", "days", 10.0, 1.0, 1.0, 2.0, 20.0, "Ongoing", "daily_progress_report_2026-08-12.pdf", "Toolbox completed")
    ],
    "2026-08-13": [
        ("CIV-PS3-FND-001", "CIV-PS3-FND-001-01", "Civil", "Pump foundation excavation and blinding", "m3", 65.0, 45.0, 20.0, 65.0, 100.0, "Complete", "daily_progress_report_2026-08-13.txt", "Blinding completed 100%"),
        ("PIP-PS3-WLD-024", "PIP-PS3-WLD-024-01", "Piping", "Field weld joints utility header", "joints", 24.0, 14.0, 6.0, 20.0, 83.3, "Ongoing", "daily_progress_report_2026-08-13.txt", "20 joints VT accepted"),
        ("PIP-PS3-HDR-100-A", "PIP-PS3-HDR-100-A-01", "Piping", "Utility header spool Section A", "m", 40.0, 15.0, 15.0, 30.0, 75.0, "Ongoing", "daily_progress_report_2026-08-13.txt", "Fit-up Section A"),
        ("PIP-PS3-SPO-015", "PIP-PS3-SPO-015-01", "Piping", "Firewater underground spool installation", "m", 80.0, 45.0, 20.0, 65.0, 81.2, "Ongoing", "daily_progress_report_2026-08-13.txt", "Spool alignment"),
        ("ELE-PS3-CT-011", "ELE-PS3-CT-011-02", "Electrical", "Place cable trench bedding at MCC-02", "m", 160.0, 90.0, 30.0, 120.0, 75.0, "Ongoing", "daily_progress_report_2026-08-13.txt", "Bedding progressing"),
        ("HSE-PS3-IND-001", "HSE-PS3-IND-001-01", "HSE", "Site safety induction and toolbox talk", "days", 10.0, 2.0, 1.0, 3.0, 30.0, "Ongoing", "daily_progress_report_2026-08-13.txt", "Toolbox completed")
    ],
    "2026-08-15": [
        ("CIV-PS3-TR-0220", "CIV-PS3-TR-0220-01", "Civil", "Utility trench excavation CH 0+220 to 0+260", "m", 40.0, 0.0, 20.0, 20.0, 50.0, "Ongoing", "daily_progress_report_2026-08-15.txt", "Trench excavation started"),
        ("CIV-PS3-FND-002", "CIV-PS3-FND-002-01", "Civil", "Pump foundation rebar and formwork", "t", 12.0, 0.0, 4.0, 4.0, 33.3, "Ongoing", "daily_progress_report_2026-08-15.txt", "Rebar tying begun"),
        ("PIP-PS3-WLD-024", "PIP-PS3-WLD-024-01", "Piping", "Field weld joints utility header", "joints", 24.0, 22.0, 2.0, 24.0, 100.0, "Complete", "daily_progress_report_2026-08-15.txt", "All 24 joints completed and VT/RT cleared"),
        ("PIP-PS3-HDR-100-A", "PIP-PS3-HDR-100-A-01", "Piping", "Utility header spool Section A", "m", 40.0, 30.0, 10.0, 40.0, 100.0, "Complete", "daily_progress_report_2026-08-15.txt", "Section A complete"),
        ("PIP-PS3-SPO-015", "PIP-PS3-SPO-015-01", "Piping", "Firewater underground spool installation", "m", 80.0, 65.0, 15.0, 80.0, 100.0, "Complete", "daily_progress_report_2026-08-15.txt", "Firewater line complete"),
        ("EQP-PS3-TK-001", "EQP-PS3-TK-001-01", "Static/Rotating Equipment", "Sump tank TK-01 internal inspection", "ea", 1.0, 0.0, 1.0, 1.0, 100.0, "Complete", "daily_progress_report_2026-08-15.txt", "Internal nozzle check accepted"),
        ("EQP-PS3-AIR-001", "EQP-PS3-AIR-001-01", "Static/Rotating Equipment", "Air receiver vessel V-103 erection", "ea", 1.0, 0.0, 1.0, 1.0, 100.0, "Complete", "daily_progress_report_2026-08-15.txt", "Vertical vessel placed on pad"),
        ("ELE-PS3-CT-011", "ELE-PS3-CT-011-02", "Electrical", "Place cable trench bedding at MCC-02", "m", 160.0, 140.0, 20.0, 160.0, 100.0, "Complete", "daily_progress_report_2026-08-15.txt", "Bedding completed 100%"),
        ("INS-PS3-JB-001", "INS-PS3-JB-001-01", "Instrumentation", "Field junction box installation", "ea", 12.0, 8.0, 4.0, 12.0, 100.0, "Complete", "daily_progress_report_2026-08-15.txt", "All 12 JBs mounted"),
        ("HSE-PS3-IND-001", "HSE-PS3-IND-001-01", "HSE", "Site safety induction and toolbox talk", "days", 10.0, 4.0, 1.0, 5.0, 50.0, "Ongoing", "daily_progress_report_2026-08-15.txt", "Toolbox completed")
    ],
    "2026-08-18": [
        ("CIV-PS3-FND-002", "CIV-PS3-FND-002-01", "Civil", "Pump foundation rebar and formwork", "t", 12.0, 4.0, 8.0, 12.0, 100.0, "Complete", "discipline_progress_2026-08-18.xlsx", "Rebar cage complete"),
        ("CIV-PS3-TR-0220", "CIV-PS3-TR-0220-01", "Civil", "Utility trench excavation CH 0+220 to 0+260", "m", 40.0, 20.0, 20.0, 40.0, 100.0, "Complete", "discipline_progress_2026-08-18.xlsx", "Trench completed 100%"),
        ("CIV-PS3-DRN-001", "CIV-PS3-DRN-001-01", "Civil", "Stormwater drainage channel installation", "m", 120.0, 60.0, 30.0, 90.0, 75.0, "Ongoing", "daily_progress_report_2026-08-18.pdf", "Precast drain laying"),
        ("PIP-PS3-HDR-100-B", "PIP-PS3-HDR-100-B-01", "Piping", "Utility header spool Section B", "m", 30.0, 0.0, 10.0, 10.0, 33.3, "Ongoing", "daily_progress_report_2026-08-18.pdf", "Section B fit-up"),
        ("ELE-PS3-CBL-001", "ELE-PS3-CBL-001-01", "Electrical", "11kV medium voltage cable pulling", "m", 500.0, 200.0, 100.0, 300.0, 60.0, "Delayed", "daily_progress_report_2026-08-18.pdf", "Planned finish date reached at 60%"),
        ("ELE-PS3-TR-005", "ELE-PS3-TR-005-01", "Electrical", "Cable tray installation Tier 2", "m", 200.0, 40.0, 50.0, 90.0, 45.0, "Ongoing", "discipline_progress_2026-08-18.xlsx", "Tray installation ongoing"),
        ("HSE-PS3-IND-001", "HSE-PS3-IND-001-01", "HSE", "Site safety induction and toolbox talk", "days", 10.0, 5.0, 1.0, 6.0, 60.0, "Ongoing", "daily_progress_report_2026-08-18.pdf", "Toolbox completed")
    ],
    "2026-08-19": [
        ("CIV-PS3-FND-003", "CIV-PS3-FND-003-01", "Civil", "Pump foundation structural concrete pour", "m3", 60.0, 0.0, 30.0, 30.0, 50.0, "Ongoing", "daily_progress_report_2026-08-19.txt", "Concrete pour 1st stage"),
        ("PIP-PS3-HDR-100-B", "PIP-PS3-HDR-100-B-01", "Piping", "Utility header spool Section B", "m", 30.0, 10.0, 8.0, 18.0, 60.0, "Ongoing", "daily_progress_report_2026-08-19.txt", "Section B spool placement (matches ROLLUP test)"),
        ("PIP-PS3-HYD-001", "PIP-PS3-HYD-001-01", "Piping", "Hydrostatic test Firewater sector 1", "test", 1.0, 0.0, 1.0, 1.0, 100.0, "Complete", "daily_progress_report_2026-08-19.txt", "Hydrotest passed at 24 bar"),
        ("EQP-PS3-PMP-101", "EQP-PS3-PMP-101-01", "Static/Rotating Equipment", "Crude pump P-101 baseplate grouting", "ea", 1.0, 0.0, 1.0, 1.0, 100.0, "Complete", "daily_progress_report_2026-08-19.txt", "Epoxy grouting complete"),
        ("INS-PS3-PT-021", "INS-PS3-PT-021-01", "Instrumentation", "Pressure transmitter mounting", "ea", 6.0, 0.0, 2.0, 2.0, 33.3, "Ongoing", "daily_progress_report_2026-08-19.txt", "PT-1021/1022 mounted"),
        ("HSE-PS3-IND-001", "HSE-PS3-IND-001-01", "HSE", "Site safety induction and toolbox talk", "days", 10.0, 6.0, 1.0, 7.0, 70.0, "Ongoing", "daily_progress_report_2026-08-19.txt", "Toolbox completed")
    ],
    "2026-08-20": [
        # EDGE-001: ELE-PS3-CBL-001 remains at 300m / 500m (60%) on 20-Aug after planned finish 18-Aug -> DELAYED
        ("ELE-PS3-CBL-001", "ELE-PS3-CBL-001-01", "Electrical", "11kV medium voltage cable pulling", "m", 500.0, 300.0, 0.0, 300.0, 60.0, "Delayed", "daily_progress_2026-08-20.csv", "Work paused awaiting cable drum delivery; planned finish 18-Aug passed"),
        # EDGE-004: In 20-Aug CSV ELE-PS3-TR-005 is reported as 120m / 200m = 60.0% (whereas 20-Aug XLSX reports 140m / 200m = 70.0%)
        ("ELE-PS3-TR-005", "ELE-PS3-TR-005-01", "Electrical", "Cable tray installation Tier 2", "m", 200.0, 90.0, 30.0, 120.0, 60.0, "Ongoing", "discipline_progress_2026-08-20.xlsx", "CSV field log claims 120m installed"),
        ("CIV-PS3-DRN-001", "CIV-PS3-DRN-001-01", "Civil", "Stormwater drainage channel installation", "m", 120.0, 90.0, 20.0, 110.0, 91.7, "Ongoing", "discipline_progress_2026-08-20.xlsx", "Precast drain installation Section A"),
        ("PIP-PS3-HDR-100-B", "PIP-PS3-HDR-100-B-01", "Piping", "Utility header spool Section B", "m", 30.0, 18.0, 7.0, 25.0, 83.3, "Ongoing", "discipline_progress_2026-08-20.xlsx", "Section B spools complete"),
        ("HSE-PS3-IND-001", "HSE-PS3-IND-001-01", "HSE", "Site safety induction and toolbox talk", "days", 10.0, 7.0, 1.0, 8.0, 80.0, "Ongoing", "daily_progress_2026-08-20.csv", "Toolbox completed")
    ],
    "2026-08-21": [
        ("CIV-PS3-FND-003", "CIV-PS3-FND-003-01", "Civil", "Pump foundation structural concrete pour", "m3", 60.0, 30.0, 30.0, 60.0, 100.0, "Complete", "daily_progress_report_2026-08-21.txt", "Structural concrete pour 100% complete"),
        ("EQP-PS3-PMP-102", "EQP-PS3-PMP-102-01", "Static/Rotating Equipment", "Crude pump P-102 laser alignment", "ea", 1.0, 0.0, 1.0, 1.0, 100.0, "Complete", "daily_progress_report_2026-08-21.txt", "Laser shaft alignment accepted"),
        ("PIP-PS3-HDR-100-C", "PIP-PS3-HDR-100-C-01", "Piping", "Utility header spool Section C", "m", 30.0, 0.0, 10.0, 10.0, 33.3, "Ongoing", "daily_progress_report_2026-08-21.txt", "Section C fit-up begun"),
        # EDGE-003: Quantity exceeds plan! Planned: 100, Cumulative: 108 ea (108%)
        ("PIP-PS3-SPT-030", "PIP-PS3-SPT-030-01", "Piping", "Pipe support secondary steel erection", "ea", 100.0, 80.0, 28.0, 108.0, 108.0, "Ongoing", "site_diary_2026-08-21.png", "Additional 8 support brackets erected per site engineer change memo"),
        ("PIP-PS3-TIE-002", "PIP-PS3-TIE-002-01", "Piping", "Tie-in spool fit-up manifold M-01", "joints", 4.0, 0.0, 2.0, 2.0, 50.0, "Ongoing", "daily_progress_report_2026-08-21.txt", "Tie-in flange alignment"),
        ("INS-PS3-CAL-003", "INS-PS3-CAL-003-01", "Instrumentation", "Control valve bench calibration", "ea", 10.0, 0.0, 4.0, 4.0, 40.0, "Ongoing", "daily_progress_report_2026-08-21.txt", "FCV-101 to 104 calibrated"),
        ("HSE-PS3-GAS-001", "HSE-PS3-GAS-001-01", "HSE", "Confined space atmospheric gas testing", "checks", 40.0, 26.0, 6.0, 32.0, 80.0, "Ongoing", "daily_progress_report_2026-08-21.txt", "Manifold pit gas tests cleared")
    ],
    "2026-08-22": [
        ("PIP-PS3-HDR-100-C", "PIP-PS3-HDR-100-C-01", "Piping", "Utility header spool Section C", "m", 30.0, 10.0, 10.0, 20.0, 66.7, "Ongoing", "field_report_2026-08-22.json", "Section C continuing"),
        ("PIP-PS3-VLV-005", "PIP-PS3-VLV-005-01", "Piping", "12-inch gate valve installation", "ea", 8.0, 5.0, 2.0, 7.0, 87.5, "Ongoing", "field_report_2026-08-22.json", "Valves torqued and tagged"),
        ("PIP-PS3-TIE-002", "PIP-PS3-TIE-002-01", "Piping", "Tie-in spool fit-up manifold M-01", "joints", 4.0, 2.0, 1.0, 3.0, 75.0, "Ongoing", "field_report_2026-08-22.json", "Manifold spool 3 joints complete"),
        ("EQP-PS3-GEN-001", "EQP-PS3-GEN-001-01", "Static/Rotating Equipment", "Diesel generator DG-01 installation", "ea", 1.0, 0.0, 1.0, 1.0, 100.0, "Complete", "field_report_2026-08-22.json", "DG-01 anchored on plinth"),
        ("INS-PS3-CAL-003", "INS-PS3-CAL-003-01", "Instrumentation", "Control valve bench calibration", "ea", 10.0, 4.0, 4.0, 8.0, 80.0, "Ongoing", "field_report_2026-08-22.json", "FCV-105 to 108 calibrated"),
        ("HSE-PS3-IND-001", "HSE-PS3-IND-001-01", "HSE", "Site safety induction and toolbox talk", "days", 10.0, 9.0, 1.0, 10.0, 100.0, "Complete", "field_report_2026-08-22.json", "10-day induction completed"),
        ("HSE-PS3-AUD-001", "HSE-PS3-AUD-001-01", "HSE", "Weekly environmental compliance audit", "audits", 2.0, 1.0, 1.0, 2.0, 100.0, "Complete", "field_report_2026-08-22.json", "Audit 2 closed out")
    ]
}

def generate_csv_reports():
    fieldnames = [
        "Report Date", "Activity ID", "Task ID", "Discipline", "Work Description",
        "Unit", "Planned Qty", "Prior Actual", "Today Actual", "Cumulative Actual",
        "Progress Pct", "Status", "Evidence Reference", "Remarks"
    ]
    for dt, rows in CSV_PROGRESS_DATA.items():
        fname = os.path.join(CSV_DIR, f"daily_progress_{dt}.csv")
        if os.path.exists(fname) and dt == "2026-08-14":
            continue

        with open(fname, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(fieldnames)
            for r in rows:
                writer.writerow([dt] + list(r))
        print(f"Generated CSV Progress Report: {fname}")

# -------------------------------------------------------------
# 5. SCANNED DIARY (21-AUG)
# -------------------------------------------------------------
def generate_scanned_diary_21aug():
    fname = os.path.join(DIARY_DIR, "site_diary_2026-08-21.png")
    if os.path.exists(fname):
        return

    width, height = 1055, 1491
    img = Image.new("RGB", (width, height), color=(250, 248, 242)) # Aged paper tone
    draw = ImageDraw.Draw(img)

    # Draw faint lined notebook ruling
    margin_left = 110
    draw.line([(margin_left, 80), (margin_left, height - 80)], fill=(220, 180, 180), width=2) # Red vertical margin
    for y in range(160, height - 80, 42):
        draw.line([(60, y), (width - 60, y)], fill=(225, 232, 240), width=1) # Blue ruling

    # Supervisor diary text entries for 21-Aug
    header_text = "SITE SUPERVISOR DAILY LOGBOOK"
    date_text = "DATE: 21-Aug-2026   SHIFT: Day (08:00 - 18:00)   LOC: Pump Station 3"
    weather_text = "WEATHER: Overcast & humid, 34 deg C. Rain held off."

    entries = [
        "1. CIV: Foundation structural pour completed for P-101/P-102 plinth.",
        "   - Batch plant sent 4 transit mixers (60 m3 total poured).",
        "   - Curing burlap placed and continuous wetting arranged.",
        "2. MECH / EQP: P-102 Pump set on baseplate. Laser shaft alignment accepted.",
        "   - Cold alignment check signed off by millwright specialist (< 0.04mm).",
        "3. PIPING: Tie-in spool fit-up at existing Manifold M-01 in progress.",
        "   - 2 field joints fitted up. 2 remaining for tomorrow's tie-in window.",
        "4. PIPING SUPPORT: Pipe support erection Secondary Steel SPT-030.",
        "   - 28 support brackets erected today (Cumulative = 108 / 100 planned).",
        "   - Site memo SM-04 approved +8 extra brackets for drainage clearance.",
        "5. INST: Control valve bench calibration at shop.",
        "   - FCV-101 through FCV-104 calibrated and tagged.",
        "6. HSE: Confined space entry permit issued for Manifold M-01 pit.",
        "   - 6 gas tests taken: O2=20.9%, LEL=0%, H2S=0ppm. Zero incidents.",
        "",
        "SIGNATURE: K. Sharma (Construction Supervisor)"
    ]

    try:
        font_header = ImageFont.truetype("arial.ttf", 26)
        font_sub = ImageFont.truetype("arial.ttf", 18)
        font_body = ImageFont.truetype("arial.ttf", 19)
    except:
        font_header = ImageFont.load_default()
        font_sub = ImageFont.load_default()
        font_body = ImageFont.load_default()

    draw.text((margin_left + 20, 95), header_text, fill=(40, 50, 70), font=font_header)
    draw.text((margin_left + 20, 130), date_text, fill=(60, 60, 80), font=font_sub)
    draw.text((margin_left + 20, 172), weather_text, fill=(20, 40, 120), font=font_body)

    cur_y = 214
    for line in entries:
        draw.text((margin_left + 20, cur_y), line, fill=(15, 30, 95), font=font_body)
        cur_y += 42

    img.save(fname, "PNG")
    print(f"Generated Scanned Diary: {fname}")

# -------------------------------------------------------------
# 6. FIELD REPORTS (JSON)
# -------------------------------------------------------------
JSON_CLAIMS = [
    {
        "file": "field_report_2026-08-12_p101_blinding.json",
        "payload": {
            "event_id": "EVT-20260812-001",
            "schedule_id": "SIH26122_NFU",
            "event_date": "2026-08-12",
            "raw_claim_text": "Blinding concrete pour continued for Crude Pump P-101/P-102 foundation base. 25 m3 placed today, total 45 m3.",
            "input_channel": "MOBILE_APP",
            "language_detected": "en",
            "reported_activity_id": "CIV-PS3-FND-001",
            "discipline": "Civil",
            "action": "POUR",
            "event_type": "PROGRESS_UPDATE",
            "claim_mode": "CUMULATIVE_QTY",
            "asset_tag": "FND-P101",
            "location": "Pump Station 3",
            "claimed_quantity": 45.0,
            "claimed_uom": "m3",
            "claimed_pct": 69.2,
            "delay_reason": None,
            "supervisor_id": "8b0821c9-6617-48f2-8957-69b596f241a1",
            "photo_path": "photos/20260812_p101_blinding.jpg",
            "status": "EXTRACTED"
        }
    },
    {
        "file": "field_report_2026-08-14_golden_claim.json",
        "payload": {
            "event_id": "EVT-20260814-001",
            "schedule_id": "SIH26122_NFU",
            "event_date": "2026-08-14",
            "raw_claim_text": "PIP-PS3-WLD-024: Two field weld joints completed on utility header today. Cumulative 22 of 24 joints accepted by VT.",
            "input_channel": "SUPERVISOR_PORTAL",
            "language_detected": "en",
            "reported_activity_id": "PIP-PS3-WLD-024",
            "discipline": "Piping",
            "action": "WELD",
            "event_type": "PROGRESS_UPDATE",
            "claim_mode": "CUMULATIVE_QTY",
            "asset_tag": "P-102",
            "location": "Pump Station 3",
            "claimed_quantity": 22.0,
            "claimed_uom": "joints",
            "claimed_pct": 91.7,
            "delay_reason": None,
            "supervisor_id": "8b0821c9-6617-48f2-8957-69b596f241a1",
            "photo_path": "photos/20260814_weld_vt_024.jpg",
            "status": "EXTRACTED"
        }
    },
    {
        "file": "field_report_2026-08-19_pump_plinth_pour.json",
        "payload": {
            "event_id": "EVT-20260819-002",
            "schedule_id": "SIH26122_NFU",
            "event_date": "2026-08-19",
            "raw_claim_text": "First stage 30 m3 structural concrete placed for pump foundation. Slump test 130mm accepted.",
            "input_channel": "MOBILE_APP",
            "language_detected": "en",
            "reported_activity_id": "CIV-PS3-FND-003",
            "discipline": "Civil",
            "action": "POUR",
            "event_type": "PROGRESS_UPDATE",
            "claim_mode": "CUMULATIVE_QTY",
            "asset_tag": "FND-P101",
            "location": "Pump Station 3",
            "claimed_quantity": 30.0,
            "claimed_uom": "m3",
            "claimed_pct": 50.0,
            "delay_reason": None,
            "supervisor_id": "3c7f12e8-9921-41b3-b452-92a101f39281",
            "photo_path": "photos/20260819_fnd_pour_stage1.jpg",
            "status": "EXTRACTED"
        }
    },
    {
        "file": "field_report_2026-08-22_generator_installation.json",
        "payload": {
            "event_id": "EVT-20260822-003",
            "schedule_id": "SIH26122_NFU",
            "event_date": "2026-08-22",
            "raw_claim_text": "Standby diesel generator DG-01 offloaded and positioned onto acoustic plinth in generator bay.",
            "input_channel": "SUPERVISOR_PORTAL",
            "language_detected": "en",
            "reported_activity_id": "EQP-PS3-GEN-001",
            "discipline": "Static/Rotating Equipment",
            "action": "ERECT",
            "event_type": "PROGRESS_UPDATE",
            "claim_mode": "CUMULATIVE_QTY",
            "asset_tag": "DG-01",
            "location": "Generator Bay",
            "claimed_quantity": 1.0,
            "claimed_uom": "ea",
            "claimed_pct": 100.0,
            "delay_reason": None,
            "supervisor_id": "3c7f12e8-9921-41b3-b452-92a101f39281",
            "photo_path": "photos/20260822_dg01_positioned.jpg",
            "status": "EXTRACTED"
        }
    }
]

def generate_field_reports():
    for item in JSON_CLAIMS:
        fname = os.path.join(JSON_DIR, item["file"])
        with open(fname, "w", encoding="utf-8") as f:
            json.dump(item["payload"], f, indent=2)
        print(f"Generated Field Claim JSON: {fname}")

if __name__ == "__main__":
    generate_txt_reports()
    generate_pdf_reports()
    generate_xlsx_reports()
    generate_csv_reports()
    generate_scanned_diary_21aug()
    generate_field_reports()
